# -*- coding: utf-8 -*-
"""Extract and normalize the raw-material list from Косметическая база.xlsx."""
import json
import re
import sys
import io
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

SRC = Path(r"F:\cosmetic-compound-base\Косметическая база.xlsx")
OUT = Path(r"F:\cosmetic-compound-base\data\raw-materials.source.json")
OUT.parent.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb["Лист1"]

rows = []
for r in range(1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    rows.append((r, a, b))


def norm(v):
    if v is None:
        return ""
    s = str(v).replace("\u00a0", " ")
    s = re.sub(r"[\t\r\n]+", " ", s)
    s = re.sub(r"\s{2,}", " ", s)
    return s.strip()


records = []
for r, a, b in rows[1:]:
    name = norm(a)
    inci = norm(b)
    if not name:
        continue
    records.append({"row": r, "name": name, "inci_source": inci})

print(f"extracted {len(records)} rows")


def slugify(text: str) -> str:
    table = {
        "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh",
        "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
        "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c",
        "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu",
        "я": "ya",
    }
    s = text.lower()
    s = "".join(table.get(ch, ch) for ch in s)
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"-{2,}", "-", s).strip("-")


seen = {}
for rec in records:
    key = re.sub(r"[^a-zа-я0-9]+", "", rec["name"].lower())
    rec["key"] = key
    if key in seen:
        seen[key]["duplicate_rows"].append(rec["row"])
        if not seen[key]["inci_source"] and rec["inci_source"]:
            seen[key]["inci_source"] = rec["inci_source"]
        continue
    rec["duplicate_rows"] = []
    rec["id"] = slugify(rec["name"])[:70]
    seen[key] = rec

unique = list(seen.values())
# ensure unique ids
used = {}
for rec in unique:
    base = rec["id"] or "item"
    if base in used:
        used[base] += 1
        rec["id"] = f"{base}-{used[base]}"
    else:
        used[base] = 1

print(f"unique {len(unique)}")

dups = [r for r in unique if r["duplicate_rows"]]
print("duplicates merged:", len(dups))
for d in dups:
    print("  ", d["name"], d["duplicate_rows"])

OUT.write_text(json.dumps(unique, ensure_ascii=False, indent=2), encoding="utf-8")
print("wrote", OUT)

lista = Path(r"F:\cosmetic-compound-base\tools\unique_list.txt")
lista.write_text(
    "\n".join(f'{i+1:3d}. {r["name"]}  |  {r["inci_source"]}' for i, r in enumerate(unique)),
    encoding="utf-8",
)
print("wrote", lista)
